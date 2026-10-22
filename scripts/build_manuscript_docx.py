"""Build a submission-style DOCX from the Markdown manuscript.

Run with the bundled Python (python-docx 1.2.0):
  /Users/mac/.dsh/dsh-runtimes/dsh-primary-runtime/dependencies/python/bin/python3 \
    scripts/build_manuscript_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "manuscript_v1.md"
OUT = ROOT / "docs" / "manuscript_v1.docx"

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(12)


def set_run_font(run, size=BODY_SIZE, bold=None, italic=None):
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text: str):
    """Add text with **bold** and *italic* markers."""
    text = text.replace("\\*", "\u0000").replace("`", "")
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2].replace("\u0000", "*"))
            set_run_font(run, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1].replace("\u0000", "*"))
            set_run_font(run, italic=True)
        else:
            run = paragraph.add_run(part.replace("\u0000", "*"))
            set_run_font(run)


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Manuscript — Journal of Electrocardiology — page ")
    set_run_font(run, size=Pt(9))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    sizes = {1: Pt(14), 2: Pt(12), 3: Pt(12)}
    run = p.add_run(text)
    set_run_font(run, size=sizes.get(level, Pt(12)), bold=True)
    if level >= 3:
        run.italic = True
    return p


def add_body(doc, text, first_indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_after = Pt(0)
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    add_inline(p, text)


def parse_table(lines):
    rows = []
    for line in lines:
        s = line.strip()
        if not s.startswith("|"):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            value = row[j] if j < len(row) else ""
            add_inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def main() -> None:
    text = SRC.read_text(encoding="utf-8")
    # Drop author notes, which must not appear in the submission file.
    text = text.split("### Notes for the author")[0]
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    add_page_number_footer(section)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    # ------------------------- title page ------------------------- #
    title = lines[0].lstrip("# ").strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(title)
    set_run_font(run, size=Pt(16), bold=True)

    # Everything between the title and "## Abstract" on the title page.
    block = []
    for line in lines[1:]:
        if line.strip().startswith("## Abstract"):
            break
        if line.strip() == "---" or not line.strip():
            continue
        block.append(line.strip())

    for line in block:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, line)
        for run in p.runs:
            run.font.size = Pt(11)

    doc.add_page_break()

    # ------------------------- body ------------------------- #
    i = 0
    # Skip to Abstract
    while i < len(lines) and not lines[i].strip().startswith("## Abstract"):
        i += 1

    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()

        if not s or s == "---":
            i += 1
            continue

        if s.startswith("### "):
            add_heading(doc, s[4:].strip(), 2)
            i += 1
            continue

        if s.startswith("## "):
            add_heading(doc, s[3:].strip(), 1)
            i += 1
            continue

        if s.startswith("# "):
            add_heading(doc, s[2:].strip(), 1)
            i += 1
            continue

        if s.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, parse_table(table_lines))
            continue

        if s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            add_inline(p, s[2:].strip())
            i += 1
            continue

        # Numbered reference entries should keep a hanging indent.
        if re.match(r"^\d+\.\s", s):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.35)
            add_inline(p, s)
            i += 1
            continue

        add_body(doc, s)
        i += 1

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
