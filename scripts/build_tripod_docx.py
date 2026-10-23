"""Build an uploadable TRIPOD+AI checklist DOCX from the Markdown mapping."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "TRIPOD_AI_checklist.md"
OUT = ROOT / "docs" / "TRIPOD_AI_checklist.docx"


def set_font(run, size=9, bold=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, text, size=9, bold=False):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold)


def main() -> None:
    lines = SRC.read_text(encoding="utf-8").splitlines()
    table_lines = [ln for ln in lines if ln.strip().startswith("|")]
    rows = []
    for ln in table_lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if all(set(c) <= set("-: ") for c in cells):
            continue
        rows.append(cells)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, margin, Inches(0.6))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "TRIPOD+AI 2024 reporting checklist", size=14, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(
        p,
        "Noise-adapted single-lead AI-ECG for hypertrophy screening: development and external validation in PTB-XL, Chapman-Shaoxing, and Georgia",
        size=10,
        bold=False,
    )
    p = doc.add_paragraph()
    add_inline(
        p,
        "Collins GS, Moons KGM, Dhiman P, et al. TRIPOD+AI statement. BMJ. 2024;385:e078378. "
        "Items not applicable to this retrospective public-data study are marked N/A with a reason.",
        size=8,
    )

    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    widths = [Inches(0.35), Inches(3.7), Inches(3.2), Inches(3.0)]
    for i, row in enumerate(rows):
        for j in range(4):
            cell = table.cell(i, j)
            cell.width = widths[j]
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            add_inline(para, row[j] if j < len(row) else "", size=8, bold=(i == 0))

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
